#!/usr/bin/env python3
"""
Re-derive the "Published spacer recovered" row of Table 1, and check the manuscript against it.

The row counts, for each tool, at how many of the seven targets the design it recommended sits
on the protospacer the source publication used. Nothing computed it: the counts were typed,
and when the PRIDICT entry at OsCDC48-T1 was corrected -- its recorded spacer did not occur in
the genome and is absent from the deposited PRIDICT run -- the typed 2 of 7 silently became
wrong. It is 3 of 7 on the corrected data, and 2 of 7 reproduces EXACTLY under the superseded
spacer, which is how the rule behind the row was recovered.

The rule is protospacer identity, nothing softer. Two conventions have to be allowed for
before comparing, because they are how the tools print, not what they designed:

  * PlantPegDesigner prints a 23 nt protospacer+PAM string, so only its first 20 nt are
    compared. The row previously read "not scored" for it, on the belief that its rendering
    shifted a base at the ends; re-running the tool on the deposited windows showed the
    strings are exact and the shifted bases were in the sequences submitted to it, so it can
    now be scored like the others.
  * PRIDICT substitutes a 5' G for Pol III transcription. A string differing from the
    published spacer ONLY at position 1, where the tool's base is G, counts as the same
    protospacer.

    python3 analysis/check_spacer_recovery.py [--docx PATH]
Exit status 1 if the manuscript and the data disagree.
"""
import sys as _s; _s.dont_write_bytecode = True   # a deposit should not ship __pycache__
import argparse, csv, os, re, sys

ap = argparse.ArgumentParser()
ap.add_argument('--docx', default=os.path.join(os.environ.get('PPE_DOCS',
    os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), 'documents')), 'Manuscript_PlantPrimeEditor.docx'))
a = ap.parse_args()


# The submitted documents are not part of the code deposit; say so and stop
# rather than crashing or reporting a missing file as a disagreement.
import sys as _sys, os as _os
_sys.path.insert(0, _os.path.join(_os.path.dirname(_os.path.abspath(__file__)), 'lib'))
from need_documents import require as _require   # noqa: E402
_require([a.docx], 'check_spacer_recovery.py', '--docx')

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)

PUBLISHED = {'OsALS-T2 (Lin 2021)': 'GGGTATGGTGGTGCAATGGG',
             'OsEPSPS-T1':          'GCAGTCACGGCTGCTGTCAA',
             'OsCDC48-T2':          'GACCAGCCAGCGTCTGGCGC',
             'OsCDC48-T1':          'GCTAGCTTTGACATAATCTC',
             'OsAAT':               'CAAGGATCCCAGCCCCGTGA',
             'OsACC-T1':            'TTCCTCGTGCTGGACAAGTG',
             'TaGW2':               'CACAAGAAAATCCACCAGGA'}


def same_protospacer(got, want):
    g = got.upper()
    if len(g) > 20:                       # protospacer+PAM string
        g = g[:20]
    if g == want:
        return True
    # a 5' G substituted for Pol III transcription, and nothing else different
    return len(g) == 20 and g[0] == 'G' and g[1:] == want[1:]


rows = list(csv.DictReader(open(os.path.join(ROOT, 'data', 'tool_comparison.csv'))))
tools = []
for r in rows:
    if r['tool'] not in tools:
        tools.append(r['tool'])
count, hits = {}, {}
for tl in tools:
    h = [t for t in PUBLISHED
         if any(same_protospacer(r['spacer'], PUBLISHED[t])
                for r in rows if r['target'] == t and r['tool'] == tl)]
    count[tl], hits[tl] = len(h), h
    print('  %-20s %d / %d   %s' % (tl, len(h), len(PUBLISHED), ', '.join(sorted(h)) or '—'))

from docx import Document
doc = Document(a.docx)
bad = []

t1 = doc.tables[0]
hdr = [c.text.strip() for c in t1.rows[0].cells]
row = next((r for r in t1.rows if r.cells[0].text.strip().startswith('Published spacer')), None)
if row is None:
    sys.exit('  FAIL  Table 1 has no "Published spacer recovered" row')
NAME = {'Plant Prime Editor': 'Plant Prime Editor', 'PlantPegDesigner': 'PlantPeg',
        'pegFinder': 'peg-Finder', 'PE-Designer': 'PE-Designer', 'PRIDICT': 'PRIDICT'}
print()
for tl in tools:
    col = next((i for i, h in enumerate(hdr) if NAME[tl] in h), None)
    if col is None:
        bad.append('Table 1 has no column for %s' % tl); continue
    cell = row.cells[col].text.strip()
    m = re.match(r'^(\d+)\s*/\s*(\d+)', cell)
    if not m:
        bad.append('Table 1 %s reads %r, not a count' % (tl, cell)); continue
    if int(m.group(1)) != count[tl] or int(m.group(2)) != len(PUBLISHED):
        bad.append('Table 1 %s reads %r, the data give %d / %d'
                   % (tl, cell, count[tl], len(PUBLISHED)))
    else:
        print('  ok    Table 1 %-20s %s' % (tl, cell))

WORD = {1: 'one', 2: 'two', 3: 'three', 4: 'four', 5: 'five', 6: 'six', 7: 'seven'}
para = next((p for p in doc.paragraphs if 'recovered it at' in p.text), None)
if para is None:
    bad.append('the Results sentence naming the recovery counts was not found')
else:
    for tl in ('PE-Designer', 'PRIDICT', 'PlantPegDesigner'):
        m = re.search(re.escape(tl) + r'[^.]{0,40}?at (\w+)', para.text)
        if not m:
            bad.append('the sentence does not give a count for %s' % tl); continue
        if m.group(1) != WORD[count[tl]]:
            bad.append('the sentence says %s recovered it at %s, the data give %s'
                       % (tl, m.group(1), WORD[count[tl]]))
        else:
            print('  ok    Results  %-20s %s' % (tl, m.group(1)))
    n = re.search(r'returned the spacer the original authors validated at (\w+) of seven',
                  para.text)
    if n and n.group(1) != WORD[count['pegFinder']]:
        bad.append('the sentence says pegFinder %s, the data give %s'
                   % (n.group(1), WORD[count['pegFinder']]))

for b in bad:
    print('  FAIL  ' + b)
print('  %s %d disagree' % ('ok   ' if not bad else 'FAIL ', len(bad)))
sys.exit(1 if bad else 0)
