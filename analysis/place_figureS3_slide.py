#!/usr/bin/env python3
"""
place_figureS3_slide.py — rebuild the Figure S3 slide around the current composite.

Figure S3 is the one main-sequence figure that has to stay a placed image: it is a set of
interface captures, so there are no shapes to draw. What it should not do is waste the slide.
The single-column composite was 2.91 x 5.65 in on a 13.33 in slide — 22 per cent of the width,
with the four screenshots too small to read and empty margins either side. With the panels laid
out two across and two down by assemble_figureS3.py the composite is close to the slide's own
proportions, so it can be placed nearly full width.

The image is scaled from its real pixel aspect rather than a hardcoded box, so recapturing a
panel at a different size cannot silently stretch the figure. The caption's resolution note is
rewritten from the file itself for the same reason.

    python3 analysis/place_figureS3_slide.py [path/to/Figures.pptx]
"""
import os, sys

from PIL import Image
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.enum.text import PP_ALIGN

from build_editable_figures import (
    DECK as DEFAULT_DECK, ROOT, IN, text, blank_slide, reorder, INK, MUTED,
)

# ── never overwrite the user's deck without being told to (added 23 Aug 2026) ──
# These scripts used to call prs.save(DECK) on the path given as argv[1], silently
# replacing a 3 MB binary the deposit does not ship and cannot restore. Confirmed
# destructive by md5 before and after a single run. The default is now to write a
# sibling file; --in-place restores the old behaviour for anyone who wants it.
def _ppe_out_path(deck_path):
    import os as _o, sys as _s
    if '--in-place' in _s.argv:
        return deck_path
    root, ext = _o.path.splitext(deck_path)
    return root + '_edited' + ext

DECK = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_DECK
IMG = os.path.join(ROOT, '..', 'FigureS3_capture_kit', 'FigureS3_interface_walkthrough.png')

SLIDE_W, SLIDE_H = 13.33, 7.50
MAX_W, TOP = 11.60, 0.78          # leaves room for the two caption lines beneath
CAP_GAP = 0.12

def _build_fingerprint():
    """Read the design-parameter fingerprint out of the shipped HTML.

    It was hardcoded, which meant the caption kept claiming a75e20c3 after the build
    had moved on. The screenshots must come from the build named here, so take the
    name from the file rather than from memory.

    The fingerprint is not a literal in the HTML - it is an FNV-1a hash the page
    computes over its own scoring source at load - so it has to be read out of a
    loaded copy of the tool rather than grepped for. The test harness already loads
    the file headlessly, so borrow it. If Node is not available, fall back to the
    build date, which is a literal.
    """
    import json, re as _re, subprocess
    html = os.path.join(ROOT, 'plant_prime_editor_v1.0.html')
    loader = os.path.join(ROOT, 'tests', 'lib', 'load_tool.js')
    if os.path.exists(loader):
        # the loader narrates each script block on stdout, so mark our own line
        js = ("const{ctx}=require(%s);"
              "process.stdout.write('\\n@@FP@@'+JSON.stringify(ctx.PPE_BUILD||{}));" % json.dumps(loader))
        try:
            out = subprocess.run(['node', '-e', js], capture_output=True, text=True,
                                 timeout=180, env=dict(os.environ, PPE_HTML=html))
            tail = (out.stdout or '').rsplit('@@FP@@', 1)
            if len(tail) == 2:
                fp = json.loads(tail[1]).get('fingerprint')
                if fp:
                    return fp
        except (OSError, ValueError, subprocess.SubprocessError):
            pass
    try:
        with open(html, encoding='utf-8', errors='replace') as fh:
            m = _re.search(r"var\s+built\s*=\s*'([^']+)'", fh.read())
        return m.group(1) if m else 'unknown'
    except OSError:
        return 'unknown'


def _caption():
    """Built at call time so the fingerprint is read from the HTML in the tree being used."""
    return ('A Module 1, allele comparison — one difference at position 306, AGG to AGA, '
            'reported as synonymous. B Module 2, assembled pegRNA — 20 nt spacer, 10 nt '
            'primer-binding site, 17 nt reverse-transcriptase template, one edit encoded. '
            'C Module 3, vector selection — the top-ranked vector for PE2 in rice. '
            'D The primer set, 19 oligos with per-primer quality control. '
            'Rice OsALS-T2, PE2, build ' + _build_fingerprint() + '.')


def main():
    if not os.path.exists(IMG):
        sys.exit('  composite not found: %s\n  run analysis/assemble_figureS3.py first' % IMG)
    im = Image.open(IMG)
    px, py = im.size
    dpi = im.info.get('dpi', (300, 300))[0] or 300
    aspect = px / py

    prs = Presentation(DECK)
    n_before = len(prs.slides._sldIdLst)
    hits = [i for i, sl in enumerate(prs.slides)
            if any(sh.has_text_frame and sh.text_frame.text.strip().startswith('Figure S3')
                   for sh in sl.shapes)]
    if len(hits) != 1:
        sys.exit('  expected exactly one Figure S3 slide, found %d' % len(hits))
    old = hits[0]

    # fit to width, then fall back to height if that would run into the caption
    w = MAX_W
    h = w / aspect
    avail_h = SLIDE_H - TOP - 0.92
    if h > avail_h:
        h = avail_h
        w = h * aspect
    x = (SLIDE_W - w) / 2.0

    s = blank_slide(prs)
    text(s, 0.42, 0.26, 12.49, 0.40, 'Figure S3   Interface walkthrough', size=17, bold=True)
    s.shapes.add_picture(IMG, IN(x), IN(TOP), IN(w), IN(h))

    cy = TOP + h + CAP_GAP
    text(s, 0.42, cy, 12.49, 0.56, _caption(), size=9.4, colour=INK)
    text(s, 0.42, cy + 0.46, 12.49, 0.20,
         'Supplied at full resolution as FigureS3_interface_walkthrough.png and .tif '
         '(%d x %d px, %d dpi = %.0f x %.0f mm). This slide is a placed image, not editable '
         'shapes.' % (px, py, dpi, 25.4 * px / dpi, 25.4 * py / dpi),
         size=8.0, italic=True, colour=MUTED)

    order = [i for i in range(n_before) if i != old]
    order.insert(old, n_before)
    reorder(prs, order, dropped=[old])
    _out = _ppe_out_path(DECK)
    prs.save(_out)
    print('  wrote ' + _out + ('  (--in-place given)' if _out == DECK else '  (original left untouched; pass --in-place to overwrite)'))

    print('  composite   : %d x %d px, %d dpi, aspect %.3f' % (px, py, dpi, aspect))
    print('  placed at   : %.2f x %.2f in  (was 2.91 x 5.65)' % (w, h))
    print('  slide width used: %.0f%%  (was 22%%)' % (100 * w / SLIDE_W))
    print('  Figure S3 rebuilt at slide %d, %d shapes' % (old + 1, len(s.shapes)))


if __name__ == '__main__':
    main()
