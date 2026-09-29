#!/usr/bin/env python3
"""
The supplementary tables exist twice — as Word tables in Supplementary_Data.docx and
as sheets in Supplementary_Tables_v1.0.xlsx — and a third time as the values the
shipped engine actually produces. Nothing checked that the three agreed.

They did not. On 13 September 2026 the two Figure 3 tables were found to carry:

  S7  an overhang pair, ACCG/CGTG, that belongs to no vector in the deposit, and a
      single pair quoted as "the" BsaI overhang when it belongs to 3 of the 12 BsaI
      vectors; and four of six rows naming the wrong primers for their route
  S11 the primer-binding-site lengths from before genPBS began ranking on the
      species melting band — OsCDC48-T1 at 11 nt where the tool returns 12, TaGW2
      at 10 where it returns 15 — and PE3b/PE5b marked Pass at a locus where the
      corrected engine declines

Both copies carried the same errors, so comparing them with each other would not
have caught it. This script checks all three sources against one another.

  python3 analysis/check_supplementary_tables.py [--docx PATH] [--xlsx PATH]

Exit status is non-zero if anything disagrees, so it can gate a release.
"""
import sys as _s; _s.dont_write_bytecode = True   # a deposit should not ship __pycache__
import argparse, csv, json, os, subprocess, sys, re

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)

ap = argparse.ArgumentParser()
ap.add_argument('--docx', default=None)
ap.add_argument('--xlsx', default=None)
a = ap.parse_args()


# The submitted documents are not part of the code deposit; say so and stop
# rather than crashing or reporting a missing file as a disagreement.
import sys as _sys, os as _os
_sys.path.insert(0, _os.path.join(_os.path.dirname(_os.path.abspath(__file__)), 'lib'))
from need_documents import require as _require   # noqa: E402
_require([a.docx, a.xlsx], 'check_supplementary_tables.py', '--docx')

def find(name, given):
    if given: return given
    for d in (ROOT, os.path.join(ROOT, '..'),
              os.environ.get('PPE_DOCS', os.path.join(ROOT, 'documents'))):
        p = os.path.join(d, name)
        if os.path.exists(p): return p
    return None

DOCX = find('Supplementary_Data.docx', a.docx)
XLSX = find('Supplementary_Tables_v1.0.xlsx', a.xlsx)
if not DOCX or not XLSX:
    print('Supplementary_Data.docx / Supplementary_Tables_v1.0.xlsx not found — nothing to check.')
    print('Pass --docx and --xlsx if they live outside the deposit.')
    sys.exit(0)

try:
    from docx import Document
    import openpyxl
except ImportError as e:
    print('python-docx and openpyxl are required:', e); sys.exit(0)

doc = Document(DOCX)
wb  = openpyxl.load_workbook(XLSX, data_only=True)
problems = []
def chk(ok, msg):
    print(('  ok   ' if ok else '  FAIL ') + msg)
    if not ok: problems.append(msg)

# ── the engine, for the values both copies are supposed to report ───────────
node = subprocess.run(['node', '-e', '''
const {ctx}=require(process.argv[1]+"/tests/probe.js"); const vm=require("vm");
console.log(vm.runInContext(`JSON.stringify(VECTORS.map(function(v){return {
  name:v.name, enzyme:v.enzyme, o5:v.overhang_5, o3:v.overhang_3,
  n5:v.nick_overhang_5, n3:v.nick_overhang_3, cassettes:v.cassettes,
  systems:v.peSystems, addgene:v.addgene||null };}))`,ctx,{timeout:30000}));
''', ROOT], capture_output=True, text=True, cwd=ROOT)
VEC = None
for line in node.stdout.splitlines():
    if line.startswith('['):
        VEC = json.loads(line); break
if VEC is None:
    print('could not load the engine; run from inside the deposit'); sys.exit(0)

# ── Table S6, keyed by vector name ──────────────────────────────────────────
print('\nTable S6 — the profiled vectors')
def s6_docx():
    ts = [t for t in doc.tables
          if [c.text.strip() for c in t.rows[0].cells][:2] == ['Vector','PE systems']]
    out = {}
    for t in ts:
        for r in t.rows[1:]:
            v = [c.text.strip() for c in r.cells]; out[v[0]] = v
    return out, [c.text.strip() for c in ts[0].rows[0].cells]
def s6_xlsx():
    ws = wb['Table S6']; out = {}; hdr = None
    for i in range(1, ws.max_row+1):
        v = [('' if ws.cell(i,j).value is None else str(ws.cell(i,j).value).strip())
             for j in range(1, ws.max_column+1)]
        if v[0] == 'Vector': hdr = v; continue
        if v[0] and v[1]: out[v[0]] = v
    return out, hdr
D6, dh = s6_docx(); X6, xh = s6_xlsx()
chk(dh == xh, 'S6 column headers identical in both copies')
chk(set(D6) == set(X6), 'S6 lists the same vectors in both copies')
diff6 = sum(1 for k in set(D6) & set(X6) for x, y in zip(D6[k], X6[k]) if x != y)
chk(diff6 == 0, 'S6 cells agree between the two copies (%d differ)' % diff6)
chk(len(D6) == len(VEC), 'S6 lists every vector the tool carries (%d vs %d)' % (len(D6), len(VEC)))
byname = {v['name']: v for v in VEC}
ALIAS = {'pICH47742::pRPS5a-PE2max-NC': 'pICH47742::pRPS5a-PE2max-NC (Vu 2024)'}
bad = []
for k, row in D6.items():
    v = byname.get(k) or byname.get(ALIAS.get(k, ''))
    if not v: bad.append(k + ': not a vector in the tool'); continue
    if row[dh.index('Enzyme')] != v['enzyme']: bad.append(k + ': enzyme')
    if row[dh.index('Cass.')] != str(v['cassettes']): bad.append(k + ': cassettes')
    m = re.search(r'#(\d+)', row[dh.index('Availability')])
    if (m.group(1) if m else None) != (v['addgene'] or None): bad.append(k + ': Addgene')
chk(not bad, 'S6 agrees with the engine' + ('' if not bad else ': ' + '; '.join(bad[:4])))

# ── Table S2, the scoring terms ─────────────────────────────────────────────
# Never checked until 13 September 2026, and the two copies had drifted: the Word table said
# the rice melting band was re-derived from "74 measurements", the workbook said 73. Both are
# real numbers from the same series — 74 recovered, 73 analysed after one 5 nt design is
# excluded, 72 plotted in Figure 4A at the 13 targets with three or more lengths — and the
# table stated one of them with no indication that the others existed. The heading also read
# "rice and other monocots" while the tool's own band carries label:'rice', evidence:'measured'
# and the table two rows down marks every other species EXTRAPOLATED.
print('\nTable S2 — the scoring terms')
def s2_docx():
    # TWO tables in the Word file carry the header ['Term','Condition'] — S1, the spacer
    # scoring terms, and S2, the primer-binding-site ones. Taking the first match compared S1
    # against the workbook's S2 sheet and reported three phantom differences while missing the
    # melting-temperature row entirely. Select on content, not on the header alone.
    for t in doc.tables:
        if [c.text.strip() for c in t.rows[0].cells][:2] != ['Term', 'Condition']: continue
        rows = {r.cells[0].text.strip(): [c.text.strip() for c in r.cells] for r in t.rows[1:]}
        if any(k.startswith('Melting temperature') for k in rows): return rows
    return {}
def s2_xlsx():
    ws = wb['Table S2']; out = {}; start = None
    for i in range(1, ws.max_row+1):
        v = [('' if ws.cell(i,j).value is None else str(ws.cell(i,j).value).strip())
             for j in range(1, 5)]
        if v[0] == 'Term': start = i; continue
        if start and v[0]: out[v[0]] = v
    return out
D2, X2 = s2_docx(), s2_xlsx()
chk(bool(D2) and bool(X2), 'S2 is present in both copies')
common = set(D2) & set(X2)
diff2 = sum(1 for k in common for x, y in zip(D2[k][:4], X2[k]) if x != y)
chk(diff2 == 0, 'S2 cells agree between the two copies (%d differ)' % diff2)

# the melting-temperature row, against the series it is derived from
tmrows = list(csv.DictReader(open(os.path.join(ROOT, 'data', 'tm_table.csv'))))
tmok = [r for r in tmrows if r.get('status') == 'OK']
_per = {}
for r in tmok: _per[r['target_id']] = _per.get(r['target_id'], 0) + 1
_ge3 = {t for t, n in _per.items() if n >= 3}
counts = {'recovered': len(tmrows), 'targets': len({r['target_id'] for r in tmrows}),
          'analysed': len(tmok),
          'plotted': sum(1 for r in tmok if r['target_id'] in _ge3), 'plot_targets': len(_ge3)}
mt = next((v for k, v in D2.items() if k.startswith('Melting temperature (SantaLucia')), None)
chk(mt is not None, 'S2 carries the measured melting-temperature row')
if mt:
    head, src = mt[0], mt[3]
    chk(head.rstrip().endswith('rice'),
        'S2 attributes the measured band to rice alone, not to monocots generally (%r)' % head[-34:])
    for label, n in (('recovered', counts['recovered']), ('analysed', counts['analysed']),
                     ('plotted', counts['plotted'])):
        chk(str(n) in src, 'S2 states the %s count, %d' % (label, n))
    chk(str(counts['targets']) in src and str(counts['plot_targets']) in src,
        'S2 states both target counts, %d and %d' % (counts['targets'], counts['plot_targets']))

# ── Table S7, the route table ───────────────────────────────────────────────
print('\nTable S7 — the assembly routes')
t7 = [t for t in doc.tables
      if [c.text.strip() for c in t.rows[0].cells][:3] == ['Strategy','Enzyme or method','Overhangs']][0]
D7 = {r.cells[0].text.strip(): [c.text.strip() for c in r.cells] for r in t7.rows[1:]}
ws = wb['Table S7']; X7 = {}; start = None
for i in range(1, ws.max_row+1):
    if ws.cell(i,1).value == 'Strategy': start = i; continue
    if start and ws.cell(i,1).value:
        X7[str(ws.cell(i,1).value).strip()] = [
            ('' if ws.cell(i,j).value is None else str(ws.cell(i,j).value).strip())
            for j in range(1, 6)]
chk(set(D7) == set(X7), 'S7 lists the same routes in both copies')
diff7 = sum(1 for k in set(D7) & set(X7) for x, y in zip(D7[k][:5], X7[k]) if x != y)
chk(diff7 == 0, 'S7 cells agree between the two copies (%d differ)' % diff7)

# a vector without a nick cassette has no nick overhangs at all
allov = set()
for v in VEC:
    for k in ('o5','o3','n5','n3'):
        if v.get(k): allov.add(v[k])
blob = ' '.join(' '.join(r) for r in D7.values())
invented = sorted({m for m in re.findall(r'\b[ACGT]{4}\b', blob)} - allov)
chk(not invented, 'every four-base overhang S7 names exists in the tool'
    + ('' if not invented else ': ' + ', '.join(invented) + ' does not'))

bsa = [v for v in VEC if v['enzyme'] == 'BsaI']
pairs = {}
for v in bsa: pairs[v['o5'] + '/' + v['o3']] = pairs.get(v['o5'] + '/' + v['o3'], 0) + 1
# the Golden Gate row, found by chemistry rather than by the enzyme it used to be named for
ggkey = next((k for k in D7 if 'Golden Gate' in k and 'swap' not in k.lower()), None)
gg = ' '.join(D7.get(ggkey, []))
chk(len(pairs) == 1 or 'read from the' in gg or 'not fixed' in gg,
    'S7 does not present one overhang pair as if it were the BsaI route\'s '
    '(%d distinct pairs across %d BsaI vectors)' % (len(pairs), len(bsa)))

# ── the route conditions, against the engine's own auto-select rule ──────────
# The rule is  bestId = publishedIsNotGG ? onestep : ggBlocked ? (bsmbiBlocked ? onestep
# : bsmbi) : gg.  Three of its four branches were missing from the table: the default route
# was named for BsaI rather than read from the acceptor, and one-step Gibson was given only
# as "Lin et al. 2021 workflow" — neither the not-published-Golden-Gate default nor the
# both-Type-IIS-routes-closed fallback.
enz = sorted({v['enzyme'] for v in VEC})
chk(ggkey is not None and 'BsaI' not in (ggkey or ''),
    'S7 does not name the default Golden Gate route for one enzyme (%s)' % (ggkey,))
chk(all(e in gg for e in enz),
    'S7\'s Golden Gate row names every enzyme the library uses (%s)' % ', '.join(enz))

gib = next((' '.join(v) for k, v in D7.items() if 'One-step' in k), '')
chk('not Golden Gate' in gib or 'Jin et al' in gib,
    'S7 gives the published-route-is-not-Golden-Gate condition for one-step Gibson')
chk('swap' in gib.lower() or 'both' in gib.lower() or 'Type IIS routes' in gib,
    'S7 gives the conflict-triggered fallback for one-step Gibson')

swap = next((' '.join(v) for k, v in D7.items() if 'swap' in k.lower()), '')
chk('acceptor' in swap or 'vector' in swap.lower(),
    'S7\'s swap row states the conflict as being with the acceptor\'s own enzyme')

# ── Table S11, the case studies ─────────────────────────────────────────────
print('\nTable S11 — the case studies')
t11 = [t for t in doc.tables
       if [c.text.strip() for c in t.rows[0].cells][:3] == ['Panel','Architecture','Locus']][0]
D11 = [[c.text.strip() for c in r.cells] for r in t11.rows[1:]]
ws = wb['Table S11']; X11 = []; start = None
for i in range(1, ws.max_row+1):
    if ws.cell(i,1).value == 'Panel': start = i; continue
    if start and ws.cell(i,1).value:
        X11.append([('' if ws.cell(i,j).value is None else str(ws.cell(i,j).value).strip())
                    for j in range(1, 11)])
chk(len(D11) == len(X11), 'S11 has the same number of rows in both copies (%d vs %d)'
    % (len(D11), len(X11)))
diff11 = sum(1 for a, b in zip(D11, X11) for x, y in zip(a, b) if x != y)
chk(diff11 == 0, 'S11 cells agree between the two copies (%d differ)' % diff11)

csj = os.path.join(HERE, 'case_studies.json')
if os.path.exists(csj):
    cs = json.load(open(csj))
    run = {}
    for panel in ('controlled', 'spread'):
        for r in cs[panel]['results']:
            g = r.get('design') or {}
            run[(r['architecture'], r['locus'])] = (
                str(g.get('pbs_len') or ''), str(g.get('rt_len') or ''),
                str(g.get('homology_beyond_edit') or ''), str(r.get('primer_count') or ''),
                'Declined' if r.get('declined') else ('Pass' if r.get('ok') else 'FAIL'))
    hdr = [c.text.strip() for c in t11.rows[0].cells]
    ix = {h: hdr.index(h) for h in ('PBS (nt)','RT (nt)','Homology (nt)','Primers','Result')}
    stale = []
    for row in D11:
        k = (row[1], row[2])
        if k not in run: stale.append('%s @ %s: not in the run' % k); continue
        want = run[k]
        got = (row[ix['PBS (nt)']], row[ix['RT (nt)']], row[ix['Homology (nt)']],
               row[ix['Primers']], row[ix['Result']])
        if got != want: stale.append('%s @ %s: %s vs %s' % (k[0], k[1], got, want))
    chk(not stale, 'S11 agrees with analysis/case_studies.json'
        + ('' if not stale else ' (%d rows differ, e.g. %s)' % (len(stale), stale[0])))
else:
    print('  note  analysis/case_studies.json absent — run node analysis/case_studies.js first')

# ── the worked example, stated in four places ───────────────────────────────
# Figure 3, Supplementary Data S1, Figure S3 and the Table S11 controlled panel all draw
# one worked example, and each was built by a script carrying its own copy of the inputs.
# On 13 September 2026 Figure 3 was found to have been rebuilt from the first OsALS-T2 row
# of the BENCHMARK — a different edit, G to T at position 301, sitting at the nick — while
# everything else used the paper's example, G to A at position 306. The figure said 14 nt
# of RT template with 13 nt of homology; the documents said 17 and 11. Both were right
# about their own input, and nothing compared them.
print('\nThe worked example — one design, four places it is stated')
wex_path = os.path.join(HERE, 'worked_example.json')
if not os.path.exists(wex_path):
    print('  note  analysis/worked_example.json absent')
else:
    W = json.load(open(wex_path))
    got = subprocess.run(['node', '-e', '''
const fs=require("fs"), vm=require("vm"), path=require("path");
const ROOT=process.argv[1];
process.env.PPE_HTML = ROOT + "/plant_prime_editor_v1.0.html";
const W=JSON.parse(fs.readFileSync(ROOT+"/analysis/worked_example.json","utf8"));
const {ctx}=require(ROOT+"/tests/probe.js");
const g=fs.readFileSync(ROOT+"/"+W.sequence_file,"utf8").split("\\n")
          .filter(l=>!l.startsWith(">")).join("").replace(/\\s+/g,"").toUpperCase();
ctx.__a={g:g, nk:W.nick_pos0, st:W.strand, sp:W.spacer,
         ed:[{genomicPos:W.edit_pos0, type:W.edit_type, ref:g.charAt(W.edit_pos0), alt:W.edit_to}]};
console.log(vm.runInContext(`JSON.stringify((function(){
  var a=__a;
  var pbs=genPBS(a.g,a.nk,a.st,${W.pbs_min},${W.pbs_max},a.sp,false);
  var rt =genRT(a.g,a.nk,a.ed,a.st,${W.rt_min},${W.rt_max},a.sp,pbs.length?pbs[0].seq:"");
  return {pbs:pbs[0].seq, pbs_len:pbs[0].seq.length, pbs_tm:pbs[0].tm,
          rt:rt[0].seq, rt_len:rt[0].seq.length,
          homology_beyond_edit:rt[0].homologyBeyondEdit};
})())`,ctx,{timeout:30000}));
''', ROOT], capture_output=True, text=True, cwd=ROOT)
    eng = None
    for line in got.stdout.splitlines():
        if line.startswith('{'): eng = json.loads(line); break
    if eng is None:
        chk(False, 'the engine returns a design for the worked example')
    else:
        E = W['expect']
        for k in ('pbs', 'pbs_len', 'rt', 'rt_len', 'homology_beyond_edit'):
            chk(eng[k] == E[k], 'worked example %s: engine %s, record %s' % (k, eng[k], E[k]))

        # Table S11's own row for the same architecture and locus must be the same design
        if os.path.exists(csj):
            row = next((r for r in cs['controlled']['results']
                        if r['architecture'] == W['pe_system'] and r['locus'] == W['locus']), None)
            d = (row or {}).get('design') or {}
            chk(bool(d), 'Table S11 carries the worked example\'s architecture at its locus')
            if d:
                chk(d.get('rt_len') == E['rt_len'] and d.get('pbs_len') == E['pbs_len']
                    and d.get('homology_beyond_edit') == E['homology_beyond_edit'],
                    'Table S11\'s %s row at %s is the same design as the figure '
                    '(PBS %s/%s, RT %s/%s, homology %s/%s)'
                    % (W['pe_system'], W['locus'], d.get('pbs_len'), E['pbs_len'],
                       d.get('rt_len'), E['rt_len'],
                       d.get('homology_beyond_edit'), E['homology_beyond_edit']))

        # and the prose that describes Supplementary Data S1
        blob = ' '.join(p.text for p in doc.paragraphs)
        for lab, val in (('reverse-transcriptase template', E['rt_len']),
                         ('primer-binding site', E['pbs_len'])):
            near = re.search(re.escape(lab) + r'[^.]{0,40}?(\d+)\s*nt', blob)
            chk(near is not None and int(near.group(1)) == val,
                'Supplementary Data states the worked example\'s %s as %d nt%s'
                % (lab, val, '' if near is None else ' (found %s)' % near.group(1)))

# ── build stamps, and the route asymmetry ──────────────────────────────────
# Two more ways the figure and the prose came apart, both found on 13 September 2026 by
# reading the rendered slide against the text.
#
#   - the slide computes its build stamp; the legend had one typed in. The moment the
#     design-parameter fingerprint moved, the slide said 5331eb12 and the legend d39dd6c8.
#   - Table S7 said a BsmBI acceptor swaps to BsaI. It does not: the only Type IIS
#     alternatives the tool offers are BsmBI and Esp3I, both CGTCTC, so on a BsmBI acceptor
#     one planted CGTCTC site closes all three Type IIS routes — which is what Figure 3A
#     shows and what the table now says.
print('\nBuild stamps and the swap direction')
fig3 = os.path.join(HERE, 'figure3_editable_data.json')
if os.path.exists(fig3):
    want = json.load(open(fig3)).get('build')
    blob = ' '.join(p.text for p in doc.paragraphs)
    gen = re.findall(r'output of build ([0-9a-f]{8}) and regenerates with '
                     r'analysis/build_figure\d+_editable\.js', blob)
    dep = re.findall(r'(?:deposited build is|against build) ([0-9a-f]{8})', blob)
    chk(all(b == want for b in gen + dep),
        'every generated-from and deposited build stamp is the current one (%s)'
        % ', '.join(sorted(set(gen + dep)) or ['none quoted']))
    # the historical capture must NOT have been restamped
    chk('572c4104' in blob,
        'the Figure S3 capture still names the build it was captured from, 572c4104')
    # ACROSS ALL THREE DOCUMENTS, and on a family of wordings rather than one string. The
    # first version of this checked only Supplementary_Data.docx for the exact phrase "differs
    # only in advisory text". Figure_Legends.docx carried "differs FROM 572c4104 only in
    # advisory text ... and finds none that differ" — the same false claim, in a document the
    # check never opened, in a wording it would not have matched anyway. It survived two full
    # verification passes.
    import glob as _glob
    EQ_BAD = re.compile(r'only in advisory text|finds none that differ')
    for _f in ('Supplementary_Data.docx', 'Manuscript_PlantPrimeEditor.docx',
               'Figure_Legends.docx'):
        _p = os.path.join(os.path.dirname(DOCX), _f)
        if not os.path.exists(_p): continue
        _d = Document(_p)
        _b = ' '.join([x.text for x in _d.paragraphs] +
                      [c.text for t in _d.tables for r in t.rows for c in r.cells])
        _hit = EQ_BAD.search(_b)
        chk(_hit is None, '%s does not claim the deposited build differs from the capture only '
            'in advisory text%s' % (_f, '' if _hit is None else ': ' + _hit.group(0)))

swap = next((' '.join(v) for k, v in D7.items() if 'swap' in k.lower()), '')
chk('Esp3I' in swap and 'CGTCTC' in swap,
    'S7\'s swap row says which enzymes the swap actually offers')
chk('BsaI acceptor' in swap and ('cannot rescue' in swap or 'shares' in swap),
    'S7\'s swap row states that the rescue works in one direction only')
# and the engine agrees: no route in the build carries a BsaI swap
node2 = subprocess.run(['node', '-e', '''
const {ctx}=require(process.argv[1]+"/tests/probe.js"); const vm=require("vm");
console.log(vm.runInContext(`JSON.stringify(Object.keys(CLONING_STRATEGY_META).map(function(k){
  return k + "=" + (CLONING_STRATEGY_META[k].recog || "none");}))`,ctx,{timeout:30000}));
''', ROOT], capture_output=True, text=True, cwd=ROOT)
routes = None
for line in node2.stdout.splitlines():
    if line.startswith('['): routes = json.loads(line); break
if routes is not None:
    iis = [r for r in routes if r.endswith('GGTCTC') or r.endswith('CGTCTC') or r.endswith('GAAGAC')]
    swaps = [r for r in iis if not r.startswith('gg=')]
    chk(all(r.endswith('CGTCTC') for r in swaps),
        'no swap route in the build reads anything but CGTCTC (%s)' % ', '.join(swaps))

# ── no data file may be shadowed at the repository root ────────────────────
# A builder that wrote its data file to the CURRENT WORKING DIRECTORY instead of to
# analysis/ left a second copy at the root whenever anyone ran it from there. On
# 13 September 2026 ./figure3_editable_data.json was found holding build c0a40a2e — the
# PRE-CORRECTION worked example, a 126 nt insert with a 14 nt reverse-transcriptase
# template — three hours staler than analysis/figure3_editable_data.json beside it, and
# shipped inside the deposit archive. Anything reading it from the root got the figure the
# corrections had just removed. Every reader and writer is anchored on __dirname now; this
# check is what stops a stray copy reappearing.
print('\nNo shadowed data files')
shadow = [f for f in os.listdir(HERE)
          if f.endswith('.json') and os.path.exists(os.path.join(ROOT, f))]
chk(not shadow,
    'no analysis/*.json is shadowed by a copy at the repository root'
    + ('' if not shadow else ': ' + ', '.join(sorted(shadow))))

# and no script names a data file relative to wherever it happens to be run from
bare = []
for d, pats in ((os.path.join(ROOT, 'analysis'), ('.js',)), (os.path.join(ROOT, 'tests'), ('.js',))):
    for f in sorted(os.listdir(d)):
        if not f.endswith(pats): continue
        src = open(os.path.join(d, f), encoding='utf-8', errors='replace').read()
        for m in re.finditer(r'(?:readFileSync|writeFileSync)\(\s*[\'"](\.{0,2}/?(?:data|analysis)/[^\'"]+|[A-Za-z0-9_]+\.json)[\'"]', src):
            bare.append('%s: %s' % (f, m.group(1)))
chk(not bare,
    'every script names its data files from __dirname, not from the working directory'
    + ('' if not bare else ': ' + '; '.join(bare[:4])))

# ── the scored benchmark must carry what the scorer adds ───────────────────
# data/benchmark_scored.csv is the output of TWO stages: merge_benchmark.py assembles the
# rows, score_batch_v2.js scores them and adds the computed columns. Running the first alone
# leaves a file that looks complete and is not: on 13 September 2026 a sweep script ran
# merge_benchmark.py on its own and the benchmark lost edit_from_top, edit_to_top and
# edit_strand_note, which took tests/audit_mnp.js from 23 passed to 21 passed, 2 failed. The
# runner caught it; this check names the cause rather than the symptom.
print('\nThe scored benchmark is fully scored')
bcsv = os.path.join(ROOT, 'data', 'benchmark_scored.csv')
if not os.path.exists(bcsv):
    chk(False, 'data/benchmark_scored.csv exists')
else:
    import csv as _csv
    with open(bcsv) as fh:
        hdr = next(_csv.reader(fh))
    need = ['edit_from_top', 'edit_to_top', 'edit_strand_note', 'rt_seq', 'composite_score',
            'published_rank_percentile']
    miss = [c for c in need if c not in hdr]
    chk(not miss, 'the scored benchmark carries every scored column'
        + ('' if not miss else ': missing ' + ', '.join(miss)
           + ' — run score_batch_v2.js over it, merge_benchmark.py alone is only half the pipeline'))

# ── Table S12, the ALS site registry, and the labels it governs ─────────────
# Two different protospacers in the rice ALS gene were both called OsALS-T2 — one by Lin 2020,
# one by Lin 2021 — and the site Lin 2021 numbers T2 is the one Lin 2020 numbers T1. A reviewer
# comparing the worked example against benchmark rows of the same name was comparing it against
# a different target. Every label now carries its study; these checks keep it that way.
print('\nTable S12 — the rice ALS target sites')
als_path = os.path.join(HERE, 'als_sites.json')
reg_path = os.path.join(HERE, 'target_sites.json')
chk(os.path.exists(reg_path), 'analysis/target_sites.json is present')
if not os.path.exists(als_path):
    chk(False, 'analysis/als_sites.json is present')
else:
    ALS = json.load(open(als_path))
    labels = [l for s in ALS['sites'] for l in s['labels']]
    chk(len(set(labels)) == len(labels), 'no label in the registry is listed twice')
    # every label the registry names must appear in the benchmark, and no bare one may survive
    bcsv = os.path.join(ROOT, 'data', 'benchmark_scored.csv')
    with open(bcsv) as fh:
        brows = list(csv.DictReader(fh))
    seen = {r['locus'] for r in brows}
    bare = sorted({l for l in seen if re.fullmatch(r'OsALS-T[12]', l)})
    chk(not bare, 'no bare OsALS-T1/T2 label survives in the benchmark'
        + ('' if not bare else ': ' + ', '.join(bare)))
    missing = [l for l in labels if l not in seen]
    chk(not missing, 'every label the registry names is in the benchmark'
        + ('' if not missing else ': ' + ', '.join(missing) + ' absent'))
    # and each label must denote exactly one spacer
    bysp = {}
    for r in brows:
        if r['locus'] in labels:
            bysp.setdefault(r['locus'], set()).add(r.get('published_spacer', ''))
    multi = {k: v for k, v in bysp.items() if len(v) > 1}
    chk(not multi, 'each ALS label denotes exactly one protospacer'
        + ('' if not multi else ': ' + '; '.join('%s -> %d' % (k, len(v)) for k, v in multi.items())))
    # the registry's own site grouping must match the spacers
    for s in ALS['sites']:
        sps = {sp for l in s['labels'] for sp in bysp.get(l, set())}
        chk(len(sps) <= 1, 'registry site %s groups labels that share one spacer (%d found)'
            % (s['site'], len(sps)))

    blob = ' '.join(p.text for p in doc.paragraphs) + ' ' + \
           ' '.join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    chk('Table S12. The two rice ALS target sites' in blob,
        'Supplementary Data carries the ALS site registry')
    for l in labels:
        chk(l in blob, 'the registry table names %s' % l)

# ── which spacer each panel is built on ─────────────────────────────────────
# The worked example's SITE appears in several figures on two different spacers: the published
# one (Figure 3, Supplementary Data S1, Table S11, Figure S1) and one of three tied at the tool's top score
# (Figure 2 panel F, Figure S3). Same site, same edit, different design — 17 nt of template in
# a 129 nt transcript against 20 nt in 132. Both legends said so in their body text and neither
# panel did, so a reader comparing the two figures saw two numbers and no reason for them.
print('\nWhich spacer each panel states')
# Both now draw the SAME design, on a spacer tied at the top score, so both must say so.
# They used to differ, and a reader comparing them met a 17 nt template against a 20 nt one.
# UPDATED 20 September 2026. These two panels used to be asserted to say "TOP-RANKED
# spacer", and that assertion was itself the error: on the published edit the worked
# example's spacer is rank 2 of 20, one of three tied at the top composite score, which is
# what analysis/worked_example_run.json records and what the Methods say. The check now
# requires the accurate phrase and forbids the superseded one.
figs = {'Figure2_ABCDEF_editable.pptx': 'tied at the tool',
        'Figure3_AB_editable.pptx':     'tied at the tool'}
try:
    from pptx import Presentation
    for fname, want in figs.items():
        p = None
        for d in (ROOT,):
            if os.path.exists(os.path.join(d, fname)): p = os.path.join(d, fname); break
        if not p:
            print('  note  %s not found beside the deposit' % fname); continue
        prs = Presentation(p)
        txt = ' '.join(sh.text_frame.text for s in prs.slides
                       for sh in s.shapes if sh.has_text_frame)
        chk(want in txt and 'top score' in txt,
            '%s states its spacer choice as one of the tied top-scoring candidates' % fname)
        chk('TOP-RANKED' not in txt,
            '%s no longer claims the single top-ranked spacer' % fname)
        # Presence of the right phrase is not absence of the wrong one. The old check asserted
        # only that "TOP-RANKED spacer" appeared; a slide could have said both.
        chk(not re.search(r'published\s+(spacer|protospacer|design)', txt, re.I),
            '%s claims no published-spacer basis anywhere on the slide' % fname)
        stale = re.findall(r'OsALS-T\d(?!\s*\(Lin)', txt)
        chk(not stale, '%s carries no unqualified ALS label (%d found)' % (fname, len(stale)))
except ImportError:
    print('  note  python-pptx not available — panel labels not checked')

# ── nothing may attach the PUBLISHED design to the worked example ───────────
# Supplementary Figure S1 legitimately draws the published spacer at this locus, and Table S12
# legitimately names it as site A's defining protospacer. Both are fine on their own. What is
# not fine is a sentence that names the published spacer and, in the same breath, the worked
# example or one of the panels that draws it — the S1 legend read "spacer GGGTATGGTGGTGCAATGGG,
# nick after g.300 ... the edit used throughout the worked example and Supplementary Data S1",
# which a reader takes as the whole design and which left two nick positions at one locus with
# no reason for them. The rule: where the published spacer and the worked example appear
# together, the worked example's own spacer must appear there too, so the contrast is stated
# rather than left to the reader. Likewise a nick position other than the record's is allowed
# only in a paragraph that has named the published spacer, i.e. one that says whose nick it is.
print('\nThe published spacer is never left attached to the worked example')
_wex = os.path.join(HERE, 'worked_example.json')
if os.path.exists(_wex):
    _W = json.load(open(_wex))
    _TOP, _NICK = _W['spacer'], _W['nick_pos0']
    _PUB = None
    _bs = os.path.join(ROOT, 'data', 'benchmark_scored.csv')
    if os.path.exists(_bs):
        for _r in csv.DictReader(open(_bs)):
            if _r.get('spacer') and _r['spacer'] != _TOP and _TOP[4:] in _r['spacer']:
                _PUB = _r['spacer']; break
    if _PUB is None:
        _PUB = 'GGGTATGGTGGTGCAATGGG'   # last resort; the CSV is the source of truth
    # the panels that draw the worked example, plus the phrase itself
    _WE = re.compile(r'worked example|Supplementary Data S1|Data S1|Table S11 controlled|'
                     r'Figure 2 panel F|Figure S3', re.I)
    _NICKRX = re.compile(r'nick(?:s|ing|ed)?\s+(?:after|at)\s+g\.(\d+)')
    for _f in ('Supplementary_Data.docx', 'Manuscript_PlantPrimeEditor.docx',
               'Figure_Legends.docx'):
        _p = os.path.join(os.path.dirname(DOCX), _f)
        if not os.path.exists(_p): continue
        _d = Document(_p)
        _paras = [x.text for x in _d.paragraphs] + \
                 [c.text for t in _d.tables for r in t.rows for c in r.cells]
        _bad = []
        for _t in _paras:
            _hasPub, _hasWE, _hasTop = _PUB in _t, bool(_WE.search(_t)), _TOP in _t
            if _hasPub and _hasWE and not _hasTop:
                _bad.append('names %s beside the worked example without naming %s'
                            % (_PUB, _TOP))
            # A nick coordinate other than the record's is licensed two ways, and only two:
            # the paragraph names the published spacer, so the number is S1's design; or the
            # paragraph also quotes the record's own nick, which is the sentence reconciling
            # the interface's label (it names the base 3' of the cut) with the prose (5').
            _ownNick = ('g.%d' % _NICK) in _t
            for _m in _NICKRX.finditer(_t):
                if int(_m.group(1)) != _NICK and not (_hasPub or _ownNick):
                    _bad.append('quotes nick g.%s with neither the published spacer nor g.%d '
                                'to own it' % (_m.group(1), _NICK))
            for _m in re.finditer(r'\bnick at g\.(\d+)', _t):
                if int(_m.group(1)) != _NICK and not (_hasPub or _ownNick):
                    _bad.append('quotes nick g.%s with neither the published spacer nor g.%d '
                                'to own it' % (_m.group(1), _NICK))
        chk(not _bad, '%s keeps the published spacer off the worked example%s'
            % (_f, '' if not _bad else ' — ' + '; '.join(sorted(set(_bad)))))

# ── no superseded copy of the worked-example record may sit beside the files ─
# A stale worked_example.json (published spacer, nick 300, 17 nt template) was shipping in the
# FINAL folder next to the documents it contradicted. The deposit's copy is the only one.
print('\nNo superseded copy of the worked-example record')
if os.path.exists(_wex):
    _canon = json.load(open(_wex))
    _stale = []
    _root = os.path.dirname(DOCX)
    for _dp, _dn, _fn in os.walk(_root):
        if 'worked_example.json' not in _fn: continue
        _q = os.path.join(_dp, 'worked_example.json')
        try:
            _o = json.load(open(_q))
        except Exception:
            _stale.append(os.path.relpath(_q, _root) + ' (unreadable)'); continue
        if (_o.get('spacer'), _o.get('nick_pos0'), _o.get('expect', {}).get('rt_len')) != \
           (_canon['spacer'], _canon['nick_pos0'], _canon['expect']['rt_len']):
            _stale.append(os.path.relpath(_q, _root))
    chk(not _stale, 'every worked_example.json beside the deliverables is the current one'
        + ('' if not _stale else ' — stale: ' + ', '.join(sorted(_stale))))

# ── no document may state a worked-example geometry the record disagrees with ───
# The earlier check regex-searched a 100-character window around the phrase
# "reverse-transcriptase template", which missed a sentence in the Figure 3 legend that still
# described the whole previous design — "a 129 nt insert of 20 nt spacer, 76 nt scaffold, 17 nt
# reverse-transcriptase template" — because the numbers sat further into a long sentence. This
# scans whole paragraphs for any insert or template length that is not the current one.
# The build-equivalence numbers the documents quote, re-measured. The advisory-only count
# moves whenever the tool's advisory text changes, which it did three times in this pass: the
# documents said 207 where build_equivalence.js reports 211. The computed-value figures are the
# load-bearing ones and did not move, but an uncorrected count beside them undermines both.
_old = os.environ.get('PPE_OLD_BUILD', '/tmp/build_572c4104.html')
if os.path.exists(_old):
    print('\nThe build-equivalence numbers the documents quote')
    _r = subprocess.run(['node', os.path.join(HERE, 'build_equivalence.js'), _old,
                         os.path.join(ROOT, 'plant_prime_editor_v1.0.html')],
                        capture_output=True, text=True, cwd=ROOT)
    _n = {}
    for _l in _r.stdout.splitlines():
        _m = re.match(r'\s*(cases|computed fields compared|'
                      r'cases differing on a computed value|'
                      r'cases differing on advisory text only)\s*:\s*(\d+)', _l)
        if _m: _n[_m.group(1)] = int(_m.group(2))
    if _n:
        for _f in ('Supplementary_Data.docx', 'Manuscript_PlantPrimeEditor.docx',
                   'Figure_Legends.docx'):
            _p = os.path.join(os.path.dirname(DOCX), _f)
            if not os.path.exists(_p): continue
            _b = ' '.join(x.text for x in Document(_p).paragraphs)
            for _lab, _key, _rx in (
                    ('computed-value cases', 'cases differing on a computed value',
                     r'(\d+) cases differ on a computed value'),
                    ('advisory-only cases', 'cases differing on advisory text only',
                     r'and (\d+) more on advisory text alone')):
                _q = [int(x) for x in re.findall(_rx, _b)]
                if not _q: continue
                chk(all(x == _n[_key] for x in _q),
                    '%s: %s stated as %s, measured %d'
                    % (_f, _lab, ', '.join(map(str, sorted(set(_q)))), _n[_key]))

print('\nNo document states a superseded worked-example geometry')
_wex = os.path.join(HERE, 'worked_example.json')
_f3  = os.path.join(HERE, 'figure3_editable_data.json')
if os.path.exists(_wex) and os.path.exists(_f3):
    _W = json.load(open(_wex)); _E = _W['expect']
    _INS = json.load(open(_f3))['panelA']['insertLen']
    _RT, _PBS = _E['rt_len'], _E['pbs_len']
    _INSERT = re.compile(r'(\d+)\s*nt insert of')
    _TEMPL  = re.compile(r'(\d+)\s*nt reverse-transcriptase template')
    for _f in ('Supplementary_Data.docx', 'Manuscript_PlantPrimeEditor.docx',
               'Figure_Legends.docx'):
        _p = os.path.join(os.path.dirname(DOCX), _f)
        if not os.path.exists(_p): continue
        _d = Document(_p)
        _paras = [x.text for x in _d.paragraphs] + \
                 [c.text for t in _d.tables for r in t.rows for c in r.cells]
        # Only where the text is TALKING ABOUT the worked example. Scanning every mention of
        # "N nt reverse-transcriptase template" flagged Table S2's citation of Anzalone 2019
        # Fig. 4b — a real 34 nt template in someone else's experiment, correctly reported.
        # A check that cries wolf on a correct citation gets switched off, which is worse than
        # not having it.
        _CTX = re.compile(r'worked example|%s' % re.escape(_W['locus']), re.I)
        bad = []
        for _t in _paras:
            if not _CTX.search(_t): continue
            for m in _INSERT.finditer(_t):
                if int(m.group(1)) != _INS: bad.append('insert %s nt' % m.group(1))
            for m in _TEMPL.finditer(_t):
                if int(m.group(1)) != _RT: bad.append('template %s nt' % m.group(1))
        chk(not bad, '%s states only the current geometry (insert %d nt, template %d nt)%s'
            % (_f, _INS, _RT, '' if not bad else ' — found ' + ', '.join(sorted(set(bad)))))

# ── the supplementary as a document, not only as a set of tables ────────────
# A full read of Supplementary_Data.docx on 14 September 2026 found eleven things no check
# looked at, because every check here was pointed at a table's CONTENTS and none at the
# document's structure: a contents list missing two of the items it indexes, a table legend
# stranded under the wrong table, Table S12 sitting past the end of the reference list, two
# citations resolving to nothing, and a Table S4 that was a different table in the two formats.
print('\nThe supplementary as a document')
_sd = os.path.join(os.path.dirname(DOCX), 'Supplementary_Data.docx')
if os.path.exists(_sd):
    _d = Document(_sd)
    _ps = [p.text for p in _d.paragraphs]
    _cells = [c.text for t in _d.tables for r in t.rows for c in r.cells]
    _blob = ' '.join(_ps) + ' ' + ' '.join(_cells)

    # everything the document defines must be in its contents list, in both formats
    _deft = {int(m) for p in _ps for m in re.findall(r'^\s*Table S(\d+)\.', p)}
    _deff = {int(m) for p in _ps for m in re.findall(r'^\s*Supplementary Figure S(\d+)\.', p)}
    _cont = [(r.cells[0].text.strip(), r.cells[1].text.strip()) for r in _d.tables[0].rows]
    _clab = ' '.join(x for x, _ in _cont)
    _missT = [n for n in sorted(_deft) if 'Table S%d' % n not in _clab and
              not re.search(r'Tables S\d+–S\d+', _clab)]
    _missT = [n for n in _missT if not (n <= 5 and 'Tables S1' in _clab)]
    _missF = [n for n in sorted(_deff) if 'Figure S%d' % n not in _clab]
    chk(not _missT and not _missF,
        'the contents table indexes every table and figure the document defines'
        + ('' if not (_missT or _missF) else ' — missing '
           + ', '.join(['Table S%d' % n for n in _missT] + ['Figure S%d' % n for n in _missF])))

    # a contents entry may not describe an item differently from the item's own title
    for _n in sorted(_deff):
        _t = next((p for p in _ps if p.strip().startswith('Supplementary Figure S%d.' % _n)), None)
        if not _t:
            continue
        _own = _t.strip()[len('Supplementary Figure S%d.' % _n):].strip().rstrip('.')
        _ent = next((c for lab, c in _cont if lab == 'Figure S%d' % _n), None)
        if _ent is None:
            continue
        chk(_ent.rstrip('.') == _own,
            'contents entry for Figure S%d matches its legend title' % _n
            + ('' if _ent.rstrip('.') == _own else ' — contents %r vs legend %r' % (_ent, _own)))

    # every table legend must sit under its own table
    _order = []
    from docx.table import Table as _T
    from docx.text.paragraph import Paragraph as _P
    for _ch in _d.element.body.iterchildren():
        if _ch.tag.endswith('}p'):
            _order.append(('p', _P(_ch, _d).text.strip()))
        elif _ch.tag.endswith('}tbl'):
            _order.append(('t', ''))
    _last_title = None
    _bad = []
    for _k, _t in _order:
        if _k == 'p':
            _m = re.match(r'^Table S(\d+)\.', _t)
            if _m:
                _last_title = int(_m.group(1))
            elif _t.startswith('Twenty-three published pegRNAs') and _last_title != 10:
                _bad.append('the Table S10 legend sits under Table S%s' % _last_title)
    chk(not _bad, 'each table legend follows its own table'
        + ('' if not _bad else ' — ' + '; '.join(_bad)))

    # nothing defined may sit after the reference list
    _idx = [i for i, (k, t) in enumerate(_order) if k == 'p' and t == 'References']
    if _idx:
        _after = [t for k, t in _order[_idx[0] + 1:]
                  if k == 'p' and re.match(r'^(Table S\d+\.|Supplementary Figure S\d+\.)', t)]
        chk(not _after, 'no table or figure is defined after the References'
            + ('' if not _after else ' — ' + ', '.join(x[:24] for x in _after)))

    # every citation in the body resolves to an entry, and every entry is cited
    if _idx:
        _reftexts = [t for k, t in _order[_idx[0] + 1:]
                     if k == 'p' and re.match(r'^[A-Z][a-zA-Z\-]+,\s+[A-Z]\.', t)]
        _body = ' '.join(t for k, t in _order[:_idx[0]] if k == 'p') + ' ' + ' '.join(_cells)
        _have = set()
        for _t in _reftexts:
            _m = re.match(r'^([A-Z][a-zA-Z\-]+),.*?\((\d{4}[a-z]?)\)', _t)
            if _m:
                _have.add((_m.group(1), _m.group(2)))
        _cited = {(x, y) for x, y in
                  re.findall(r'\b([A-Z][a-zA-Z\-]+)(?:\s+[A-Z]{1,3})?\s+(?:et al\.?|and [A-Z][a-zA-Z\-]+),?\s+\(?(\d{4})\)?', _body)}
        _dangling = sorted({(x, y) for x, y in _cited if x in {n for n, _ in _have} | {
            'Anderson', 'Jin', 'Doench', 'Hsu', 'Ma', 'Dang', 'SantaLucia', 'Anzalone',
            'Lin', 'Zong', 'Li', 'Vu', 'Zhao', 'Ni', 'Nelson', 'Chen', 'Liu'}
            and (x, y) not in _have})
        chk(not _dangling, 'every citation in the supplementary resolves to a reference entry'
            + ('' if not _dangling else ' — ' + ', '.join('%s %s' % p for p in _dangling)))
        _uncited = sorted({p for p in _have
                           if not re.search(re.escape(p[0]) + r'[^\n]{0,45}?' + p[1], _body)})
        chk(not _uncited, 'every reference entry is cited in the supplementary'
            + ('' if not _uncited else ' — ' + ', '.join('%s %s' % p for p in _uncited)))

    # the ranked cohort is counted in SITES, not in labels
    _bs = os.path.join(ROOT, 'data', 'benchmark_scored.csv')
    if os.path.exists(_bs):
        _b = list(csv.DictReader(open(_bs)))
        _rk = [r for r in _b if str(r.get('published_rank_percentile', '')).strip()]
        # LABELS are not SITES. Four pairs of labels denote one protospacer and five single
        # labels denote two, so the site must be keyed on the PROTOSPACER
        # (analysis/target_sites.json). Collapsing only the ALS pair gave 25 here by luck — a
        # missed merge and a missed split cancelling — while giving 33 where the 162-passing
        # stage has 35 sites and 26 where the 141-scored stage has 27.
        import sys as _sys
        _sys.path.insert(0, os.path.join(HERE, 'lib'))
        from target_site import site_of_row as _site

        _ns = len({_site(r) for r in _rk})
        _m = re.search(r'(\d+) at (\d+) target sites entered the design-recovery ranking', _blob)
        chk(bool(_m) and int(_m.group(1)) == len(_rk) and int(_m.group(2)) == _ns,
            'the design-recovery cohort is stated as %d pegRNAs at %d sites%s'
            % (len(_rk), _ns, '' if not _m else ' — document says %s at %s'
               % (_m.group(1), _m.group(2))))

    # Table S4 must be the weights the engine applies, in both formats
    _s4 = None
    for _t in _d.tables:
        if [c.text.strip() for c in _t.rows[0].cells][:1] == ['Component']:
            _s4 = _t
            break
    if _s4 is not None:
        _w = re.findall(r'(\d+)%', ' '.join(c.text for r in _s4.rows for c in r.cells))
        _src = open(os.path.join(ROOT, 'plant_prime_editor_v1.0.html'), encoding='utf8').read()
        _rs = re.search(r'const rawSpec = \(([^;]+)\);', _src)
        _eng = [str(int(round(float(x) * 100))) for x in re.findall(r'\*\s*(0\.\d+)', _rs.group(1))] \
            if _rs else []
        chk(bool(_eng) and [x for x in _w if x in _eng] and set(_eng) <= set(_w),
            'Table S4 carries the intrinsic-risk weights the engine applies (%s)'
            % '/'.join(_eng) + ('' if set(_eng) <= set(_w) else ' — document has %s' % '/'.join(_w)))

# ── the end-to-end narrative, against the run it describes ──────────────────
# The "Worked example: from allele to ordered oligonucleotides" section and the Results
# paragraph introducing it were maintained by hand and described a different design from the
# one that produced Supplementary Data S1 and Supplementary Figure S3, while saying they were
# the same run. Part had been refreshed against the engine and part had not. Every length and
# temperature the section states must now be one the run actually returns.
_run = os.path.join(HERE, 'worked_example_run.json')
if os.path.exists(_run) and os.path.exists(_sd):
    print('\nThe end-to-end narrative against the run it describes')
    _R = json.load(open(_run))
    _lens = {len(_R['spacer']['seq']), _R['pbs'][0]['len'], _R['pbs'][1]['len'],
             _R['pbs'][2]['len'], _R['rt']['len'], _R['rt']['runner_up_len'],
             _R['insert_nt'], _R['scaffold_nt'], _R['polyT_nt'], _R['window_nt'],
             _R['rt']['homology_beyond_edit'], _R['spacer']['nick_to_edit_nt'],
             5}                       # 5 nt is the stated minimum homology the flap needs
    _lens |= {p['len'] for p in _R['primers']}
    # the P3 fold label carries its own lengths ("6 bp stem (5 G:C), 3 nt loop")
    _lens |= {int(x) for x in re.findall(r'(\d+)', _R['named']['p3fold'])}
    _wj = json.load(open(os.path.join(HERE, 'worked_example.json')))
    _lens |= {_wj['pbs_min'], _wj['pbs_max'], _wj['rt_min'], _wj['rt_max']}
    # The narrative also states the CLONED fragment and the ordered gBlock, which are
    # longer than the transcript: Module 3 prepends the guanine the U6 promoter
    # initiates at, and the synthesis order adds the enzyme flanks. Both are generated
    # values, so take them from the generated protocol rather than allowing any number.
    _prot = os.path.join(ROOT, 'DataS1', 'S1d_bench_protocol.txt')
    if os.path.exists(_prot):
        _pt = open(_prot).read()
        _lens |= {int(x) for x in re.findall(r'(\d+) nt total \(insert', _pt)}
        _lens |= {int(x) for x in re.findall(r'insert (\d+) nt \+ RE flanks', _pt)}
    _temps = {str(_R['named']['dTm']), '52', '72', '14', '20'}
    _temps |= {str(p['tm']) for p in _R['primers']}
    _temps |= {str(_R['pbs'][i]['tm_nn']) for i in range(3)}
    _temps |= {str(_R['pbs'][0]['tm_wallace'])}
    _temps = {t.rstrip('0').rstrip('.') if '.' in t else t for t in _temps} | _temps

    _sec = [t for t in [p.text for p in Document(_sd).paragraphs]
            if t.startswith(('To show that the platform delivers', 'Module 2 returned',
                             'Module 3 ranked'))]
    chk(len(_sec) == 3, 'the worked-example section has its three module paragraphs (%d found)'
        % len(_sec))
    _bad = []
    for _t in _sec:
        for _m in re.finditer(r'(\d+)\s*nt\b', _t):
            if int(_m.group(1)) not in _lens:
                _bad.append('%s nt' % _m.group(1))
        for _m in re.finditer(r'(\d+(?:\.\d+)?)\s*°C', _t):
            _v = _m.group(1)
            if _v not in _temps and _v.rstrip('0').rstrip('.') not in _temps:
                _bad.append('%s °C' % _v)
    chk(not _bad, 'every length and temperature in the narrative is one the run returns'
        + ('' if not _bad else ' — not in the run: ' + ', '.join(sorted(set(_bad)))))

    _need = [_R['spacer']['seq'], _R['locus'], _R['vector']['enzyme'],
             _R['vector']['overhang_5'], _R['vector']['overhang_3'],
             _R['codon']['ref'] + ' to ' + _R['codon']['alt'], _R['build']]
    _abs = [x for x in _need if x not in ' '.join(_sec)]
    chk(not _abs, 'the narrative names the run’s spacer, locus, enzyme, overhangs, codon '
        'change and build' + ('' if not _abs else ' — missing ' + ', '.join(_abs)))

    # and it may not name a spacer the run does not use
    _other = [s for s in re.findall(r'\b[ACGT]{20}\b', ' '.join(_sec))
              if s != _R['spacer']['seq']]
    chk(not _other, 'the narrative names no other 20 nt spacer'
        + ('' if not _other else ' — ' + ', '.join(sorted(set(_other)))))

# ── the display items are exactly the ones the paper claims ────────────────
# Supplementary Figure S5 was withdrawn on 14 September 2026. The set is Figures 1-6 and
# Supplementary Figures S1-S4, and it is asserted here rather than left to whatever happens to
# be in the folder, so neither a withdrawn figure reappearing nor a kept one going missing can
# pass unnoticed. S5 had been appended rather than inserted, so nothing renumbered when it left.
MAIN_FIGS = ['1', '2', '3', '4', '5', '6']
SUPP_FIGS = ['S1', 'S2', 'S3', 'S4']
print('\nThe display items')
if os.path.exists(_sd):
    _d2 = Document(_sd)
    _defined = sorted({m for p in _d2.paragraphs
                       for m in re.findall(r'^\s*Supplementary Figure (S\d+)\.', p.text)},
                      key=lambda x: int(x[1:]))
    chk(_defined == SUPP_FIGS,
        'the supplementary defines exactly %s' % ', '.join(SUPP_FIGS)
        + ('' if _defined == SUPP_FIGS else ' — found ' + ', '.join(_defined)))

_figdir = os.path.dirname(DOCX)
_have = os.listdir(_figdir)
# Figure S3 is a capture of the live interface, so it ships as an image and has no editable
# deck by design; every other display item must have one.
_IMAGE_ONLY = {'S3': ('.png', '.tif')}
_missing = []
for _n in MAIN_FIGS + SUPP_FIGS:
    _pre, _exts = 'Figure%s_' % _n, _IMAGE_ONLY.get(_n, ('.pptx',))
    if not any(f.startswith(_pre) and f.endswith(_exts) for f in _have):
        _missing.append('Figure %s (%s)' % (_n, ' or '.join(_exts)))
chk(not _missing, 'every display item has a file beside the documents'
    + ('' if not _missing else ' — none for ' + ', '.join(_missing)))
_keep_pre = tuple('Figure%s_' % n for n in MAIN_FIGS + SUPP_FIGS)
_stray = sorted({f for f in _have if re.match(r'Figure(S\d+|\d+)_.*\.pptx$', f)
                 and not f.startswith(_keep_pre)})
chk(not _stray, 'no editable deck belongs to a figure the paper does not carry'
    + ('' if not _stray else ' — ' + ', '.join(_stray)))

try:
    from pptx import Presentation as _Pr
    for _deck, _howto in (('Figures_PlantPrimeEditor.pptx', True),
                          ('Figures_PlantPrimeEditor_SUBMISSION.pptx', False)):
        _p = os.path.join(_figdir, _deck)
        if not os.path.exists(_p):
            continue
        _prs = _Pr(_p)
        _txt = ' '.join(sh.text_frame.text for s in _prs.slides
                        for sh in s.shapes if sh.has_text_frame)
        _found = sorted({m for m in re.findall(r'Figure (S\d+)\b', _txt)},
                        key=lambda x: int(x[1:]))
        _bad = [f for f in _found if f not in SUPP_FIGS]
        chk(not _bad, '%s carries no withdrawn figure' % _deck
            + ('' if not _bad else ' — ' + ', '.join(_bad)))
        _n = len(_prs.slides._sldIdLst)
        chk(_n == (15 if _howto else 14),
            '%s has %d slides' % (_deck, 15 if _howto else 14)
            + ('' if _n == (15 if _howto else 14) else ' — found %d' % _n))
except ImportError:
    print('  note  python-pptx not available — decks not checked')

# ── the manuscript against its own figures and data ────────────────────────
# The Results and the Limitations both said the leave-one-out test put the best primer-binding
# site "inside the 8-11 nt window in 13 of 13 cases", which data/loto.json records with
# out_of_sample:false and a note saying the window was SELECTED to achieve exactly that. The
# figure's own legend had been corrected to say so; the manuscript that cites the figure had
# not, and it also quoted 11 of 11 rankings where the data give 10, and P = 5e-4 where the data
# give 9.8e-4. Nothing compared the paper with the figure it points at.
_ms = os.path.join(os.path.dirname(DOCX), 'Manuscript_PlantPrimeEditor.docx')
_lj = os.path.join(ROOT, 'data', 'loto.json')
if os.path.exists(_ms) and os.path.exists(_lj):
    print('\nThe manuscript against data/loto.json')
    _L = json.load(open(_lj))
    _T, _C = _L['tests'], _L['cohort']
    _md = Document(_ms)
    _mps = [p.text.strip() for p in _md.paragraphs]
    _mall = ' '.join(_mps)

    # the ranking cohort, wherever the manuscript states it
    _r = _C['targets_in_ranking_test']
    _wrong = re.findall(r'positive at (?:all )?(\d+)(?: of (\d+))? targets', _mall)
    _bad = [x for x in _wrong if int(x[0]) != _r or (x[1] and int(x[1]) != _r)]
    chk(not _bad, 'the manuscript states the ranking cohort as %d targets' % _r
        + ('' if not _bad else ' — found ' + ', '.join(x[0] for x in _bad)))

    # the leave-one-out counts it may quote
    _n = _C['targets_in_argmax_test']
    _kwin = _T['inside_leave_one_out_window']['k']
    for _need, _what in ((r'%d of %d' % (_kwin, _n), 'the leave-one-out window count %d of %d'
                          % (_kwin, _n)),
                         (r'within %d nt' % _T['within_2nt_of_leave_one_out_optimum']['max_deviation_nt'],
                          'the leave-one-out deviation')):
        chk(bool(re.search(_need, _mall)), 'the manuscript states %s' % _what)

    # and it may not present the fitted count as an out-of-sample test
    _ctx = [t for t in _mps if 'leave-one-target-out' in t or 'Refitting without a target' in t]
    _claim = [t[:60] for t in _ctx
              if re.search(r'inside the [\d–-]+ nt window in %d of %d' % (_n, _n), t)]
    chk(not _claim, 'no paragraph presents the fitted %d of %d count as the out-of-sample test'
        % (_n, _n) + ('' if not _claim else ' — ' + '; '.join(_claim)))

    # the sign-test P, to one significant figure
    _p = _T['ranking']['sign_test_P_one_sided']
    _e = 0
    _x = _p
    while _x < 1:
        _x *= 10
        _e -= 1
    _m = round(_x)
    if _m >= 10:
        _m, _e = _m // 10, _e + 1
    _sup = ''.join({'-': '⁻', '0': '⁰', '1': '¹', '2': '²', '3': '³',
                    '4': '⁴', '5': '⁵', '6': '⁶', '7': '⁷', '8': '⁸',
                    '9': '⁹'}[c] for c in str(_e))
    _want = '%d × 10%s' % (_m, _sup)
    _quoted = re.findall(r'sign test P = ([^;.)]+)', _mall)
    chk(all(q.strip() == _want for q in _quoted) and _quoted,
        'every sign-test P in the manuscript reads %s' % _want
        + ('' if all(q.strip() == _want for q in _quoted) and _quoted
           else ' — found ' + ', '.join(sorted({q.strip() for q in _quoted}))))

    # the supplemental-information pointer must match what the supplementary defines
    _sd2 = Document(_sd)
    _t = sorted({int(m) for x in _sd2.paragraphs for m in re.findall(r'^\s*Table S(\d+)\.', x.text)})
    _f = sorted({int(m) for x in _sd2.paragraphs
                 for m in re.findall(r'^\s*Supplementary Figure S(\d+)\.', x.text)})
    _want_si = ('Supplementary Tables S%d–S%d, Supplementary Figures S%d–S%d and '
                'Supplementary Data S1 are available online.' % (min(_t), max(_t), min(_f), max(_f)))
    _got = next((t for t in _mps if t.startswith('Supplementary Tables S')), '')
    chk(_got == _want_si, 'the supplemental-information line points at what the supplementary '
        'defines' + ('' if _got == _want_si else ' — says %r' % _got[:80]))

    # every figure the manuscript legends must have one title, shared with the legend file
    _fl = os.path.join(os.path.dirname(DOCX), 'Figure_Legends.docx')
    if os.path.exists(_fl):
        _RX = r'^((?:Supplementary )?Figure S?\d+)\.\s*(.*)$'
        def _titles(path):
            return {m.group(1): m.group(2)
                    for x in Document(path).paragraphs
                    for m in [re.match(_RX, x.text.strip())] if m}
        _a, _b, _c = _titles(_ms), _titles(_fl), _titles(_sd)
        _dis = sorted({k for k in set(_a) | set(_b) | set(_c)
                       if len({v for v in (_a.get(k), _b.get(k), _c.get(k)) if v}) > 1})
        chk(not _dis, 'every figure has one title across the three documents'
            + ('' if not _dis else ' — disagree on ' + ', '.join(_dis)))

    # the stated word count, measured
    _names = [x.text.strip() for x in _md.paragraphs]
    if 'Abstract' in _names and 'Materials availability' in _names:
        _lo, _hi = _names.index('Abstract'), _names.index('Materials availability')
        _w = sum(len(x.text.split()) for x in list(_md.paragraphs)[_lo:_hi]
                 if x.text.strip() and not x.style.name.startswith('Heading')
                 and not x.text.strip().startswith('Keywords:'))
        _m2 = re.search(r'Word count, main text: ([\d,]+)', _mall)
        chk(bool(_m2) and int(_m2.group(1).replace(',', '')) == _w,
            'the stated word count is the measured one (%s)' % format(_w, ',')
            + ('' if _m2 and int(_m2.group(1).replace(',', '')) == _w
               else ' — states %s' % (_m2.group(1) if _m2 else 'nothing')))

# ── the case-study panel, against analysis/case_studies.json ────────────────
# The manuscript said "Forty-two of the 44 passed" and named PE3b and PE5b at the controlled
# locus as declined. It went to 44 of 44 on the belief that the worked example's own edit
# licensed a conditional nicking guide where the benchmark's edit did not.
#
# FIX DECLINE-TALLY (19 September 2026). Two things were wrong with that, and this block
# could not see either.
#
#   * The belief is false. genNickSgRNA returns ZERO PE3b-licensed guides at the controlled
#     locus under the worked example's nick AND under the published one, while PE3 — which
#     chooses by distance alone — returns eight. Both architectures decline at that locus
#     whichever edit is installed.
#   * The tally counted a decline as "a result with no checks array". A decline has a FULL
#     checks array: its seventh entry is "conditional nick declined, correctly", pass true.
#     So the count was structurally always zero, and the line "0 declined" was printed
#     beside two declines for as long as the block existed.
#
# 44 of 44 pass is correct and unchanged — a correct decline is a pass. What was wrong is
# "none declined", and the tally is now taken from the `declined` field, cross-checked
# against the Result column of Table S11 and against what the documents say.
_cs = os.path.join(HERE, 'case_studies.json')
if os.path.exists(_cs) and os.path.exists(_ms):
    print('\nThe case-study panel against analysis/case_studies.json')
    _C = json.load(open(_cs))

    def _tally(key):
        res = _C[key]['results']
        dec = [r for r in res if r.get('declined')]
        ok = [r for r in res if r.get('checks') and r.get('ok')]
        bad = [r for r in res if not (r.get('checks') and r.get('ok'))]
        return len(res), len(ok), len(dec), len(bad)

    _nC, _okC, _decC, _badC = _tally('controlled')
    _nS, _okS, _decS, _badS = _tally('spread')
    _tot, _ok, _dec = _nC + _nS, _okC + _okS, _decC + _decS
    _mt = ' '.join(x.text for x in Document(_ms).paragraphs)
    _st = ' '.join(x.text for x in Document(_sd).paragraphs) + ' ' + \
          ' '.join(c.text for t in Document(_sd).tables for r in t.rows for c in r.cells)
    for _f, _b in (('manuscript', _mt), ('supplementary', _st)):
        _claim = re.search(r'(?:All )?(\w+[-\w]*) of the (\d+)(?: pass| designs| that)', _b)
        _has = ('%d of %d' % (_ok, _tot)) in _b or ('All %d' % _tot) in _b
        chk(_has, '%s states the panel as %d of %d' % (_f, _ok, _tot)
            + ('' if _has else ' — found %r' % (_claim.group(0) if _claim else 'nothing')))
        _stale = re.search(r'Forty-two of the 44|42 of the 44|42 of 44', _b)
        chk(not _stale, '%s carries no superseded panel count' % _f
            + ('' if not _stale else ' — %r' % _stale.group(0)))
        # a document may not say nothing declined while the run declines
        _nodec = re.search(r'none declined|no design .{0,40}declined', _b)
        chk(not (_dec and _nodec), '%s does not claim nothing declined (%d do)' % (_f, _dec)
            + ('' if not (_dec and _nodec) else ' — %r' % _nodec.group(0)))
        # nor that a conditional guide exists at a locus where the engine finds none.
        # Two phrasings, because the two documents said it differently: the supplementary
        # "a conditional nicking guide exists for PE3b and PE5b", the manuscript "At this
        # edit one does." Matching only the first left the manuscript's copy uncaught.
        _exists = re.search(r'conditional nicking guide exists|At this edit one does', _b)
        chk(not (_dec and _exists),
            '%s does not claim a conditional nicking guide exists where the engine declines'
            % _f + ('' if not (_dec and _exists) else ' — %r' % _exists.group(0)))
    chk(_badC + _badS == 0,
        'every case-study combination passed its conditions (%d of %d)' % (_ok, _tot))
    chk(_dec == 2 and _decS == 0,
        'the declines are the two conditional architectures at the controlled locus '
        '(%d declined: %s)' % (_dec, ', '.join(
            r['architecture'] for k in ('controlled', 'spread')
            for r in _C[k]['results'] if r.get('declined')) or 'none'))
    # and Table S11's Result column must carry exactly those declines
    _rescol = [r[-1] for r in D11]
    chk(_rescol.count('Declined') == _dec,
        'Table S11 marks %d row(s) Declined, matching the run (%d)'
        % (_rescol.count('Declined'), _dec))

print('\n%s' % ('all checks pass' if not problems else '%d PROBLEM(S)' % len(problems)))
sys.exit(1 if problems else 0)
