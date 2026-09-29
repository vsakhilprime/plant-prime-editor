#!/usr/bin/env python3
"""
Re-derive the two Spearman coefficients the Results quote, and check the manuscript against them.

The sentence is the most direct support the melting-temperature recalibration has, and it was
the one number in the paper nothing checked. It is computed over a cohort that exists in NO
single file: the five tomato rows in data/benchmark_scored.csv carry a measured efficiency,
and the 44 digitised Lin 2020 efficiencies live only in data/benchmark_lin2020_joined.csv.
Reading either alone gives n = 5 or n = 44, not 49, and recomputing from the wrong one makes
the published coefficients look unreproducible when they are not.

The comparison is proximity -- negative absolute deviation -- of each design's primer-binding
site melting temperature to two targets: the species band (16 C for rice and tomato, 30 C for
the Triticeae and maize) and the single 30 C figure the plant literature quotes on the
Wallace scale. analysis/analyse_benchmark.py section 4 is the computation; this re-implements
it independently so a change in either would surface as a disagreement rather than be copied.

    python3 analysis/check_benchmark_correlations.py [--docx PATH]
Exit status 1 if the manuscript and the data disagree.
"""
import sys as _s; _s.dont_write_bytecode = True   # a deposit should not ship __pycache__
import argparse, csv, os, re, sys

ap = argparse.ArgumentParser()
ap.add_argument('--docx', default=os.path.join(os.environ.get('PPE_DOCS',
    os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), 'documents')), 'Manuscript_PlantPrimeEditor.docx'))
a = ap.parse_args()


# The submitted documents are not part of the code deposit; say so and stop
# rather than crashing or reporting a missing file as a disagreement.
import sys as _sys, os as _os
_sys.path.insert(0, _os.path.join(_os.path.dirname(_os.path.abspath(__file__)), 'lib'))
from need_documents import require as _require   # noqa: E402
_require([a.docx], 'check_benchmark_correlations.py', '--docx')

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
from scipy import stats

BAND = {'rice': 16, 'default': 16, 'triticeae_maize': 30}

bench = list(csv.DictReader(open(os.path.join(ROOT, 'data', 'benchmark_scored.csv'))))
join = {r['id']: r['measured_efficiency']
        for r in csv.DictReader(open(os.path.join(ROOT, 'data',
                                                  'benchmark_lin2020_joined.csv')))}
num = lambda v: float(v) if (v or '').strip() not in ('', 'n.r.', 'NA') else None

tm, eff, tgt = [], [], []
own = 0
for r in bench:
    e = num(r.get('measured_efficiency'))
    if e is None:
        e = num(join.get(r['id']))
    else:
        own += 1
    t = num(r.get('pbs_tm'))
    if e is None or t is None:
        continue
    tm.append(t); eff.append(e); tgt.append(BAND.get(r.get('pbs_tm_band') or 'default', 16))

n = len(eff)
r_band, p_band = stats.spearmanr([-abs(t - g) for t, g in zip(tm, tgt)], eff)
r_30, p_30 = stats.spearmanr([-abs(t - 30) for t in tm], eff)
print('  cohort   %d designs (%d in benchmark_scored.csv, %d joined from Lin 2020)'
      % (n, own, n - own))
print('  species band   rho = %+.3f   P = %.3g' % (r_band, p_band))
print('  30 C (Wallace) rho = %+.3f   P = %.3g' % (r_30, p_30))

from docx import Document
doc = Document(a.docx)
text = '\n'.join(p.text for p in doc.paragraphs)
para = next((p for p in doc.paragraphs if 'Spearman ρ = +' in p.text and 'whereas' in p.text), None)
if para is None:
    sys.exit('  FAIL  the Results sentence carrying both coefficients was not found')

bad = []
def want(label, pat, value, fmt):
    m = re.search(pat, para.text)
    if not m:
        bad.append('%s: the sentence does not state a value' % label); return
    got = m.group(1)
    exp = fmt % value
    if got != exp:
        bad.append('%s: the manuscript says %s, the data give %s' % (label, got, exp))
    else:
        print('  ok    %-24s %s' % (label, got))

want('n', r'Across the ([0-9]+) pegRNAs', n, '%d')
want('species-band rho', r'ρ = \+([0-9.]+), P', r_band, '%.3f')
want('species-band P', r'ρ = \+[0-9.]+, P = ([0-9.]+)\)', p_band, '%.4f')
want('30 C rho', r'ρ = −([0-9.]+), P', abs(r_30), '%.3f')
want('30 C P', r'ρ = −[0-9.]+, P = ([0-9.]+)\)', p_30, '%.5f')
if 'three species' in para.text:
    sp = {r['species'] for r in bench if num(r.get('measured_efficiency')) is not None
          or num(join.get(r['id'])) is not None}
    print('  ok    species in the cohort    %d (%s)' % (len(sp), ', '.join(sorted(sp))))
    if len(sp) != 3:
        bad.append('the sentence says three species, the cohort has %d' % len(sp))

for b in bad:
    print('  FAIL  ' + b)
print('  %s %d check(s) disagree' % ('ok   ' if not bad else 'FAIL ', len(bad)))
sys.exit(1 if bad else 0)
